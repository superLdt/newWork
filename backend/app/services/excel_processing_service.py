# -*- coding: utf-8 -*-
"""
Excel文件处理服务
提供Excel文件的解析、生成和处理功能
"""

import pandas as pd
from io import BytesIO
from typing import Dict, List, Any, Tuple, Optional
import json
from .vehicle_validation_service import VehicleValidationService


class ExcelProcessingService:
    """Excel文件处理服务类"""
    
    # 导入模板字段映射（中文字段名 -> 数据库字段名）
    FIELD_MAPPING = {
        '车牌号': 'license_plate',
        '车厢号': 'carriage_number', 
        '容积': 'actual_volume',
        '车辆类型': 'vehicle_type',
        '车辆分类': 'vehicle_category',
        '常用公司': 'frequent_companies',
        '备注': 'notes',
        '供应商类型': 'supplier_type',
        '状态': 'status'
    }
    
    # 反向映射（数据库字段名 -> 中文字段名）
    REVERSE_FIELD_MAPPING = {v: k for k, v in FIELD_MAPPING.items()}
    
    # 必填字段
    REQUIRED_FIELDS = ['actual_volume']  # 容积必填
    
    # 车辆类型选项
    VEHICLE_TYPE_OPTIONS = ['5吨', '8吨', '12吨', '20吨', '30吨', '40吨A', '40吨B']
    
    # 车辆分类选项
    VEHICLE_CATEGORY_OPTIONS = ['单车', '挂车']
    
    # 状态选项
    STATUS_OPTIONS = ['待确认', '已确认', '使用中', '维修中', '停用']
    
    @classmethod
    def create_import_template(cls) -> bytes:
        """
        创建车辆信息导入模板Excel文件
        
        Returns:
            bytes: Excel文件的二进制数据
        """
        # 创建示例数据
        template_data = {
            '车牌号': ['皖A12345', '京B67890', '沪C11111'],
            '车厢号': ['皖A36T3挂', '', '沪C22222挂'],
            '容积': [25.5, 30.0, 18.8],
            '车辆类型': ['5吨', '8吨', '5吨'],
            '车辆分类': ['挂车', '单车', '挂车'],
            '常用公司': ['公司A,公司B', '公司C', '公司D,公司E,公司F'],
            '备注': ['测试数据1', '测试数据2', '测试数据3'],
            '供应商类型': ['委办公司', '班组', '承运商'],
            '状态': ['待确认', '已确认', '待确认']
        }
        
        df = pd.DataFrame(template_data)
        
        # 创建Excel文件
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # 写入数据到工作表
            df.to_excel(writer, sheet_name='车辆信息', index=False, startrow=3)
            
            # 获取工作表对象
            worksheet = writer.sheets['车辆信息']
            
            # 添加说明信息
            worksheet['A1'] = '车辆信息导入模板'
            worksheet['A2'] = '说明：1. 车牌号或车厢号至少填写一个；2. 容积必填；3. 车厢号格式如：皖A36T3挂；4. 常用公司用逗号分隔'
            
            # 设置列宽
            column_widths = {
                'A': 15,  # 车牌号
                'B': 15,  # 车厢号
                'C': 10,  # 容积
                'D': 12,  # 车辆类型
                'E': 12,  # 车辆分类
                'F': 25,  # 常用公司
                'G': 20,  # 备注
                'H': 15,  # 供应商类型
                'I': 12   # 状态
            }
            
            for col, width in column_widths.items():
                worksheet.column_dimensions[col].width = width
            
            # 设置标题行样式
            from openpyxl.styles import Font, PatternFill, Alignment
            
            title_font = Font(bold=True, size=14)
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color='E6F3FF', end_color='E6F3FF', fill_type='solid')
            center_alignment = Alignment(horizontal='center', vertical='center')
            
            # 设置标题样式
            worksheet['A1'].font = title_font
            worksheet['A1'].alignment = center_alignment
            
            # 合并标题单元格
            worksheet.merge_cells('A1:I1')
            worksheet.merge_cells('A2:I2')
            
            # 设置表头样式
            for col in range(1, 10):  # A到I列
                cell = worksheet.cell(row=4, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_alignment
        
        output.seek(0)
        return output.getvalue()
    
    @classmethod
    def parse_excel_file(cls, file_content: bytes, filename: str = '') -> Dict[str, Any]:
        """
        解析Excel文件内容
        
        Args:
            file_content: Excel文件的二进制内容
            filename: 文件名（用于错误提示）
            
        Returns:
            Dict[str, Any]: 解析结果，包含success状态和data或error信息
        """
        try:
            # 读取Excel文件
            df = pd.read_excel(BytesIO(file_content), sheet_name=0, skiprows=3)
            
            # 检查是否为空文件
            if df.empty:
                return {
                    'success': False,
                    'error': '文件内容为空，请检查Excel文件是否包含数据'
                }
            
            # 检查必要的列是否存在
            required_columns = ['车牌号', '车厢号', '容积']
            missing_columns = [col for col in required_columns if col not in df.columns]
            if missing_columns:
                return {
                    'success': False,
                    'error': f'缺少必要的列：{", ".join(missing_columns)}'
                }
            
            # 转换字段名
            df_renamed = df.rename(columns=cls.FIELD_MAPPING)
            
            # 数据清洗
            df_cleaned = cls._clean_dataframe(df_renamed)
            
            # 转换为字典列表
            records = df_cleaned.to_dict('records')
            
            # 过滤掉完全空白的行
            filtered_records = []
            for record in records:
                # 检查是否所有关键字段都为空
                key_fields = ['license_plate', 'carriage_number', 'actual_volume']
                if any(cls._is_not_empty(record.get(field)) for field in key_fields):
                    filtered_records.append(record)
            
            if not filtered_records:
                return {
                    'success': False,
                    'error': '文件中没有有效的数据行'
                }
            
            return {
                'success': True,
                'data': filtered_records,
                'total_rows': len(filtered_records)
            }
            
        except Exception as e:
            return {
                'success': False,
                'error': f'文件解析失败：{str(e)}'
            }
    
    @classmethod
    def _clean_dataframe(cls, df: pd.DataFrame) -> pd.DataFrame:
        """
        清洗DataFrame数据
        
        Args:
            df: 原始DataFrame
            
        Returns:
            pd.DataFrame: 清洗后的DataFrame
        """
        # 创建副本避免修改原数据
        df_clean = df.copy()
        
        # 处理NaN值
        df_clean = df_clean.fillna('')
        
        # 处理字符串字段，去除首尾空格
        string_fields = ['license_plate', 'carriage_number', 'vehicle_type', 
                        'vehicle_category', 'frequent_companies', 'notes', 
                        'supplier_type', 'status']
        
        for field in string_fields:
            if field in df_clean.columns:
                df_clean[field] = df_clean[field].astype(str).str.strip()
                # 将字符串'nan'转换为空字符串
                df_clean[field] = df_clean[field].replace('nan', '')
        
        # 处理数值字段
        if 'actual_volume' in df_clean.columns:
            df_clean['actual_volume'] = pd.to_numeric(df_clean['actual_volume'], errors='coerce')
        
        return df_clean
    
    @classmethod
    def _is_not_empty(cls, value: Any) -> bool:
        """
        检查值是否不为空
        
        Args:
            value: 要检查的值
            
        Returns:
            bool: 值是否不为空
        """
        if value is None:
            return False
        if pd.isna(value):
            return False
        if isinstance(value, str) and value.strip() == '':
            return False
        if isinstance(value, (int, float)) and value == 0:
            return False
        return True
    
    @classmethod
    def validate_and_process_data(cls, records: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        验证和处理导入数据
        
        Args:
            records: 原始数据记录列表
            
        Returns:
            Dict[str, Any]: 验证和处理结果
        """
        # 使用验证服务进行批量验证
        validation_results = VehicleValidationService.validate_batch_data(records)
        
        # 获取验证摘要
        summary = VehicleValidationService.get_validation_summary(validation_results)
        
        return {
            'validation_results': validation_results,
            'summary': summary,
            'preview_data': validation_results[:10]  # 只返回前10条作为预览
        }
    
    @classmethod
    def export_validation_errors(cls, validation_results: List[Dict[str, Any]]) -> bytes:
        """
        导出验证错误报告为Excel文件
        
        Args:
            validation_results: 验证结果列表
            
        Returns:
            bytes: Excel文件的二进制数据
        """
        # 筛选出有错误的记录
        error_records = [result for result in validation_results if not result['valid']]
        
        if not error_records:
            # 如果没有错误，创建一个空的报告
            df = pd.DataFrame({'说明': ['所有数据验证通过，无错误记录']})
        else:
            # 创建错误报告数据
            report_data = []
            for result in error_records:
                row_data = {
                    '行号': result['row'],
                    '车牌号': result['original_data'].get('license_plate', ''),
                    '车厢号': result['original_data'].get('carriage_number', ''),
                    '容积': result['original_data'].get('actual_volume', ''),
                    '错误信息': '; '.join(result['errors'])
                }
                report_data.append(row_data)
            
            df = pd.DataFrame(report_data)
        
        # 创建Excel文件
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='验证错误报告', index=False)
            
            # 设置列宽
            worksheet = writer.sheets['验证错误报告']
            worksheet.column_dimensions['A'].width = 8   # 行号
            worksheet.column_dimensions['B'].width = 15  # 车牌号
            worksheet.column_dimensions['C'].width = 15  # 车厢号
            worksheet.column_dimensions['D'].width = 10  # 容积
            worksheet.column_dimensions['E'].width = 50  # 错误信息
        
        output.seek(0)
        return output.getvalue()
    
    @classmethod
    def get_field_mapping_info(cls) -> Dict[str, Any]:
        """
        获取字段映射信息
        
        Returns:
            Dict[str, Any]: 字段映射和选项信息
        """
        return {
            'field_mapping': cls.FIELD_MAPPING,
            'reverse_mapping': cls.REVERSE_FIELD_MAPPING,
            'required_fields': cls.REQUIRED_FIELDS,
            'vehicle_types': cls.VEHICLE_TYPE_OPTIONS,
            'vehicle_categories': cls.VEHICLE_CATEGORY_OPTIONS,
            'status_options': cls.STATUS_OPTIONS
        }
    
    @classmethod
    def validate_file_format(cls, filename: str, file_size: int, max_size: int = 10 * 1024 * 1024) -> Dict[str, Any]:
        """
        验证文件格式和大小
        
        Args:
            filename: 文件名
            file_size: 文件大小（字节）
            max_size: 最大允许文件大小（字节），默认10MB
            
        Returns:
            Dict[str, Any]: 验证结果
        """
        errors = []
        
        # 检查文件扩展名
        if not filename.lower().endswith(('.xlsx', '.xls')):
            errors.append('文件格式不支持，请上传Excel文件（.xlsx或.xls格式）')
        
        # 检查文件大小
        if file_size > max_size:
            errors.append(f'文件大小超过限制，最大允许{max_size // (1024*1024)}MB')
        
        if file_size == 0:
            errors.append('文件为空，请选择有效的Excel文件')
        
        return {
            'valid': len(errors) == 0,
            'errors': errors
        }